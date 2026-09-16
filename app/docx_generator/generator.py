from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.schemas.meeting import MeetingExtraction


class ProtocolDocxGenerator:
    def __init__(
            self,
            template_path: str | Path | None = None,
    ):
        if template_path is None:
            template_path = (Path(__file__).parents[1]
                             / 'templates'
                             / 'protocol_template.docx')

        self.template_path = Path(template_path)

    # Сохранение в папку проекта
    def generate(
            self,
            extraction: MeetingExtraction,
            output_path: str | Path,
    ) -> Path:
        output_path = Path(output_path)

        output_path.parent.mkdir(
            parents=True, exist_ok=True
        )

        doc = self._build_document(extraction)

        doc.save(output_path)

        return output_path

    # Выстраиваем документ
    def _build_document(
            self,
            extraction: MeetingExtraction,
    ):
        doc = Document(self.template_path)

        heard_heading = self._find_paragraph(
            doc,
            'Заслушали:',
        )

        end_marker = self._find_paragraph_starting_with(
            doc,
            '[При необходимости добавить',
        )

        self._remove_placeholder_blocks(
            doc,
            start_paragraph=heard_heading,
            end_paragraph=end_marker,
        )

        for block in extraction.heard:
            self._insert_heard_block(
                anchor=end_marker,
                block=block,
            )

        self._remove_paragraph(end_marker)

        return doc

    #Метод для поиска нужной строки
    @staticmethod
    def _find_paragraph(doc, text: str):
        for paragraph in doc.paragraphs:
            if paragraph.text.strip() == text:
                return paragraph

        raise ValueError(
            f"Paragraph was not found in {text}"
        )

    # Альтернативный поиск строки, который начинается с определённого текста
    @staticmethod
    def _find_paragraph_starting_with(doc, text: str):
        for paragraph in doc.paragraphs:
            if paragraph.text.strip().startswith(text):
                return paragraph

        raise ValueError(
            f"Paragraph was not found in {text}"
        )

    # Функция для удаления текста заглушек
    def _remove_placeholder_blocks(
            self,
            doc,
            start_paragraph,
            end_paragraph,
    ):
        paragraphs = doc.paragraphs

        start_index = next(
            i
            for i, paragraph in enumerate(paragraphs)
            if paragraph._element is start_paragraph._element
        )

        end_index = next(
            i
            for i, paragraph in enumerate(paragraphs)
            if paragraph._element is end_paragraph._element
        )

        for paragraph in paragraphs[
            start_index + 1:end_index
        ]:
            self._remove_paragraph(paragraph)

    # Функция, удаляющая строку
    @staticmethod
    def _remove_paragraph(paragraph):
        element = paragraph._element
        parent = element.getparent()

        parent.remove(element)

    # Применение стилей для сгенерированного текста
    @staticmethod
    def _add_run(
            paragraph,
            text: str,
            *,
            bold: bool = False,
            italic: bool = False,
    ):
        run = paragraph.add_run(text)

        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

        return run

    # Вставка текста
    def _insert_heard_block(
            self,
            anchor,
            block,
    ):
        speaker_paragraph = (
            anchor.insert_paragraph_before()
        )

        self._add_run(
            speaker_paragraph,
            self._speaker_heading(block),
            bold=True,
        )

        summary_paragraph = (
            anchor.insert_paragraph_before()
        )

        summary_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
        )

        self._add_run(
            summary_paragraph,
            block.summary,
        )

        if not block.resolutions:
            return

        decision_heading = (
            anchor.insert_paragraph_before()
        )

        self._add_run(
            decision_heading,
            'Решение:',
            bold=True,
        )

        for resolution in block.resolutions:
            resolution_paragraph = (
                anchor.insert_paragraph_before()
            )

            self._add_run(
                resolution_paragraph,
                self._format_resolution(resolution),
            )

    # Вставка оглавления (Спикера/Докладчика)
    @classmethod
    def _speaker_heading(cls, block) -> str:
        if block.speaker_name:
            return block.speaker_name

        return cls._format_speaker_id(
            block.speaker_id
        )

    # Форматирование Спикера, если не нашёл имя
    @staticmethod
    def _format_speaker_id(speaker_id: str) -> str:
        try:
            number = int(
                speaker_id.rsplit('_', 1)[1]
            )

            return f'Спикер {number + 1}'

        except (ValueError, IndexError):
            return 'Спикер'

    # Форматирование решений
    @staticmethod
    def _format_resolution(resolution) -> str:
        parts = [
            resolution.text.rstrip('.')
        ]

        if resolution.responsible_name:
            parts.append(
                f'Ответственный — '
                f'{resolution.responsible_name}'
            )

        if resolution.deadline:
            parts.append(
                f'срок — {resolution.deadline}'
            )

        return '— ' + '; '.join(parts) + '.'